"""Build the requested Word evidence document from real completed artifacts."""
import csv
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent
ART = ROOT / 'artifacts'
EVIDENCE = ROOT.parent / 'cloud-ml-evidence'
DEST = ROOT.parent / 'deliverables'


def csv_rows(name):
    with (ART / name).open(newline='', encoding='utf-8-sig') as stream:
        return list(csv.DictReader(stream))


def num(value, digits=3):
    if value is None or value == '':
        return 'Not available'
    try:
        return f'{float(value):.{digits}f}'
    except (ValueError, TypeError):
        return str(value)


def main():
    DEST.mkdir(exist_ok=True)
    config = json.loads((ROOT / 'config.json').read_text())
    reactive = json.loads((ART / 'live_reactive_summary.json').read_text())
    predictive = json.loads((ART / 'live_predictive_summary.json').read_text())
    def first_event(policy, field):
        rows = csv_rows(f'live_{policy}_control.csv')
        return next((float(r['elapsed_s']) / 60 for r in rows if float(r[field]) >= 2), None)
    def failure_types(policy):
        rows = csv_rows(f'live_{policy}_requests.csv')
        return {
            'transport': sum(r['reason'].startswith('timeout_or_connection_error') for r in rows),
            'routing': sum(r['reason'] == 'no_healthy_worker' for r in rows),
            'backlog': sum(r['reason'] == 'generator_backlog' for r in rows),
        }
    cleanup = json.loads((EVIDENCE / 'cleanup-verification.json').read_text())
    if not reactive['complete'] or not predictive['complete']:
        raise RuntimeError('Both real experiments must be complete before building the report.')
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.6)
    section.left_margin = section.right_margin = Inches(.65)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        style = document.styles[name]
        style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.line_spacing = 1.05
    document.styles['Title'].font.size = Pt(22)
    document.styles['Heading 1'].font.size = Pt(16)
    document.styles['Heading 2'].font.size = Pt(12)
    document.styles['Caption'].font.size = Pt(10)
    document.core_properties.author = 'Parth Tendulkar'
    document.core_properties.title = 'Cloud Resource Management Implementation Evidence'

    def para(text, style=None):
        return document.add_paragraph(text, style)

    def heading(text, new_page=False):
        if new_page:
            document.add_page_break()
        document.add_heading(text, level=1)

    def table(headers, rows, widths=None):
        t = document.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        if widths is None:
            widths = [7.2 / len(headers)] * len(headers)
        for i, width in enumerate(widths):
            t.columns[i].width = Inches(width)
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = str(h)
        for row in rows:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = str(value)
        for r, row in enumerate(t.rows):
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement('w:cantSplit'); tr_pr.append(cant_split)
            if r == 0:
                repeat = OxmlElement('w:tblHeader'); tr_pr.append(repeat)
            for c, cell in enumerate(row.cells):
                cell.width = Inches(widths[c]); cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ('top', 'left', 'bottom', 'right'):
                    line = OxmlElement(f'w:{edge}')
                    for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                        line.set(qn(f'w:{k}'), v)
                    borders.append(line)
                props.append(borders)
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    margin = OxmlElement(f'w:{edge}'); margin.set(qn('w:w'), '90'); margin.set(qn('w:type'), 'dxa')
                    margins.append(margin)
                props.append(margins)
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'DCE6F1' if r == 0 else ('F5F7FA' if r % 2 == 0 else 'FFFFFF'))
                props.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)
                    for run in p.runs:
                        run.font.size = Pt(10)
                        run.bold = r == 0
        para('')

    figure_counter = 0

    def figure(path, caption, max_height=3.8):
        nonlocal figure_counter
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f'Required real evidence missing: {path}')
        with Image.open(path) as image:
            width, height = image.size
        inches = min(7.2, max_height * width / height)
        p = document.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(inches))
        figure_counter += 1
        para(f'Figure {figure_counter}. {caption}', 'Caption')

    document.add_paragraph('Cloud Resource Management Implementation Evidence', 'Title')
    para('Parth Tendulkar | Third year Computer Engineering')
    para('Implementation date ' + datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime('%d %B %Y'))
    para('This project predicts HTTP workload five minutes ahead and starts or stops a second real AWS EC2 worker. '
         'The following pages record the implementation, measured tests and resource cleanup. '
         'Training and load generation ran on the laptop; only the application workers ran in AWS.')
    table(['Item', 'Implemented value'], [
        ['Cloud region', 'Sydney ap-southeast-2; Mumbai was blocked by the account service control policy'],
        ['Resources', 'Two Linux t3.micro workers; one 8 GiB gp3 root volume each; Standard CPU credits'],
        ['Algorithm', 'Random Forest; 120 trees; maximum depth 10; fixed random seed'],
        ['Forecast and control', 'Five-minute forecast; minute sampling; approximately ten-second live safety checks'],
        ['Account coverage', 'AWS Free account plan retained; initial displayed credit balance USD 100'],
    ], [1.65, 5.55])
    figure(EVIDENCE / '01-starting-credits.jpg', 'Starting AWS credit balance before deployment. Billing estimates are delayed.', 2.7)

    heading('Cloud deployment', True)
    para('The application was installed automatically using EC2 user data and a systemd service. '
         'No load balancer, NAT Gateway, SageMaker service or Elastic IP was created. '
         'Only TCP port 5000 from the laptop public IP was permitted; no SSH key was created.')
    figure(EVIDENCE / '04-two-running-workers.jpg', 'The two micro workers deployed in the permitted Sydney region.', 3.4)
    figure(EVIDENCE / '05-application-health-crop.jpg', 'The real application health endpoint returned status ok.', .8)
    para('Each work request burns approximately 30 ms of thread CPU time. A laptop controller reads health and CPU endpoints, '
         'discovers the new public address after a restart and routes requests only to application-healthy workers.')
    para('Controller permissions were limited to reading EC2 state and starting or stopping Project CloudML workers. '
         'The access key was encrypted for the current Windows user, excluded from screenshots, and removed during cleanup.')

    heading('Workload and prediction accuracy', True)
    para('Input data comprises eight days of generated minute-level HTTP request rates: a repeating 20-minute cycle, '
         'daily modulation and seeded noise. This is a synthetic workload, not a NASA or production trace. '
         'Features use the latest completed minute, nine earlier lags, trailing five- and fifteen-minute means '
         'and standard deviations, daily sine and cosine terms, and minute of hour.')
    para('Six days train the model, one day validates it and one day tests it. Labels crossing split boundaries are '
         'excluded. Every model predicts the same request rate five minutes ahead. Repeated synthetic patterns make '
         'this easier than an unpredictable real workload.')
    accuracy = [r for r in csv_rows('accuracy.csv') if r['split'] == 'test']
    table(['Model', 'MAE RPS', 'RMSE RPS', 'WAPE percent'],
          [[r['model'], num(r['MAE']), num(r['RMSE']), num(r['WAPE_pct'])] for r in accuracy], [2.7, 1.4, 1.4, 1.7])
    figure(ART / 'comparison.png', 'Held-out simulation. Forecasts are aligned with their future target times. Latency is a proxy.', 4.0)

    heading('Capacity and scaling mechanism', True)
    para(f'Measured demo capacity is {config["capacity_rps"]:.3f} requests per second per worker. '
         'The passing burst rate was scaled to the t3.micro 10 percent CPU baseline with a 20 percent reserve, '
         'reduced further when confirmation failed, and then verified for sixty seconds before either policy test. '
         'This is a conservative working capacity, not maximum throughput. '
         f'The standby took {config["boot_seconds"]:.1f} seconds from restart request to an application-healthy response. '
         'This calibration is not a guarantee of sustained capacity on a burstable instance.')
    calibration = csv_rows('calibration.csv')
    table(['Rate RPS', 'CPU percent', 'p95 ms', 'SLA failures percent', 'Pass'],
          [[r['rate_rps'], num(r['cpu_percent'],1), num(r['p95_success_latency_ms'],1),
            num(r['sla_violation_pct'],2), r['passes']] for r in calibration], [1.1, 1.3, 1.4, 2.2, 1.2])
    para('Predictive desired workers equals the forecast multiplied by 1.20, divided by measured capacity and rounded '
         'up, with a minimum of one and maximum of two. Reactive control adds the standby when observed demand exceeds '
         '80 percent of ready capacity or CPU exceeds 80 percent. Predictive fallback uses 90 percent demand or 80 percent CPU.')
    para('Scale-in requires three low-demand minutes, CPU below 35 percent and at least three minutes of standby readiness. '
         'The predictive policy must also forecast one worker. The standby is drained before stopping; pending starts are not duplicated.')
    figure(EVIDENCE / '06-calibration.jpg', 'Calibration results recorded from actual EC2 requests and CPU samples.', 2.4)

    heading('Live scaling demonstration', True)
    para('Both policies replayed the same fifteen-minute arrival-rate schedule, including a surprise spike and a sharp drop. '
         'Arrivals were scheduled independently of response latency. Each run began with one worker and the same historical seed; '
         'the predictive model received only past observed arrivals. These screenshots show the real instance state changes.')
    para(f'The predictive controller first requested two workers at minute {num(first_event("predictive", "desired"),2)}, '
         f'compared with minute {num(first_event("reactive", "desired"),2)} for reactive control. '
         'This demonstrates anticipation; the cost and latency tables determine whether that earlier allocation was beneficial.')
    figure(EVIDENCE / '07-reactive-scaling.jpg', 'Reactive experiment showing the standby requested in response to observed demand.', 3.4)
    figure(EVIDENCE / '08-predictive-scaling.jpg', 'Predictive experiment showing actual worker allocation while the ML controller runs.', 3.4)

    heading('Measured live results', True)
    table(['Metric', 'Reactive', 'Predictive'], [
        ['Attempted requests', str(reactive['attempted']), str(predictive['attempted'])],
        ['Successful request p95 ms', num(reactive['p95_success_latency_ms'],1), num(predictive['p95_success_latency_ms'],1)],
        ['Request errors percent', num(reactive['error_pct']), num(predictive['error_pct'])],
        ['SLA violations percent', num(reactive['sla_violation_pct']), num(predictive['sla_violation_pct'])],
        ['Mean measured CPU percent', num(reactive['avg_measured_cpu_pct']), num(predictive['avg_measured_cpu_pct'])],
        ['Sampled workload instance hours', num(reactive['workload_instance_hours'],4), num(predictive['workload_instance_hours'],4)],
        ['Modeled workload compute USD', num(reactive['compute_cost_model_usd'],5), num(predictive['compute_cost_model_usd'],5)],
        ['Modeled energy kWh', num(reactive['energy_model_kwh'],5), num(predictive['energy_model_kwh'],5)],
        ['Generator backlog count', str(reactive['generator_backlog']), str(predictive['generator_backlog'])],
    ], [3.6, 1.8, 1.8])
    para('The SLA counts failed requests, timeouts and requests over 1000 ms once each. '
         'The p95 value excludes unsuccessful requests, so it must be read with the error and SLA columns. '
         'Instance-hours and compute estimates cover sampled workload time, not setup, calibration, preflight or final shutdown.')
    rf, pf = failure_types('reactive'), failure_types('predictive')
    para('In this paired live test, prediction started the standby earlier but did not improve latency or the SLA. '
         'Predictive control consumed more instance time and modeled energy. The two policies recorded the same number '
         'of request errors, and network effects limit conclusions from a single short pair of runs.')
    para(f'Failure classification: reactive recorded {rf["transport"]} transport errors and {rf["routing"]} requests '
         f'without a healthy route; predictive recorded {pf["transport"]} and {pf["routing"]}, respectively. '
         'These include connection and health-routing effects and must not all be attributed to CPU overload or forecast error.')
    figure(ART / 'live_comparison.png', 'Actual ready-worker counts, successful-request latency and sampled CPU for both runs.', 2.2)

    heading('Prediction errors and resource tradeoffs', True)
    para('An unexpected spike can exceed the forecast. Demand and CPU fallback request the second worker, but startup '
         'still takes time and the two-worker maximum may remain insufficient. After a sudden drop, overprediction leaves '
         'spare capacity until conservative scale-in. Live forecasts update from past observations; the separate simulation '
         'fault tests deliberately keep the original forecast fixed while altering actual demand.')
    sims = csv_rows('simulation_results.csv')
    table(['Scenario', 'Policy', 'Overloaded minutes percent', 'Instance hours'],
          [[r['scenario'].replace('_',' '), r['policy'], num(r['overloaded_minutes_pct']), num(r['instance_hours'])] for r in sims],
          [2.1, 1.2, 2.3, 1.6])
    figure(ART / 'live_forecast.png', 'Live forecasts plotted at their target times rather than incorrectly compared with current demand.', 2.6)
    para('Simulation utilization is demand divided by ready capacity, capped at one; excess demand is reported separately. '
         'The simulated latency proxy is 30 divided by one minus utilization, capped at utilization 0.98. '
         'Its percentile of minute-level values is not a measured request-level p95 or SLA rate.')

    heading('CPU monitoring cost and energy', True)
    figure(EVIDENCE / '09-cloudwatch-cpu.jpg', 'AWS basic CPU monitoring for the application worker over the experiment period.', 4.8)
    para(f'Compute is modeled at the Linux On-Demand rate displayed by AWS, USD {config["compute_hourly_usd"]:.4f} '
         'per worker-hour. Full gross cost also includes EBS storage, public IPv4 address-hours and any billable transfer. '
         'Credits cover gross consumption; they do not make that consumption physically free. The final payable amount '
         'must be read from Billing after delayed usage posts.')
    para('Energy is illustrative rather than measured: assume 10 W idle and 25 W at full CPU for each worker and '
         'interpolate linearly. Sum power multiplied by interval hours and divide by 1000 to obtain kWh. '
         'Startup is assigned idle power; missing CPU samples may underestimate energy. Halving or doubling assumed '
         'power halves or doubles the energy estimate.')
    para('The long-run simulation uses modeled workload utilization for power, while live energy uses sampled CPU. '
         'Neither method measures physical server power.')
    para('Limitations include synthetic repeated traffic, one pair of short live runs, laptop routing and network latency, '
         'burstable CPU credits, changing cloud startup time, and a maximum of two workers. CPU-credit throttling can '
         'change effective capacity even when workload prediction is accurate. Benefits are judged from the actual results, '
         'not assumed in advance.')

    heading('Resource cleanup', True)
    para('Traffic and controllers were stopped before cleanup. The two demo instances were terminated, and attached '
         'root volumes and remaining project resources were checked. Termination is necessary because merely stopping '
         'instances leaves EBS storage allocated.')
    figure(EVIDENCE / '10-terminated-workers.jpg', 'Both project workers after termination.', 3.3)
    figure(EVIDENCE / '11-no-project-volumes.jpg', 'Storage verification after instance termination.', 3.3)

    heading('Final account and access verification', True)
    para(cleanup.get('summary', 'Cleanup verification is recorded in the supporting inventory file.'))
    figure(EVIDENCE / '12-final-billing.jpg', 'Billing and credit status shown after cleanup. Estimates can take about 24 hours to update.', 3.2)
    figure(EVIDENCE / '13-controller-removed.jpg', 'The temporary controller identity and its access key removed after the experiment.', 2.8)
    para('The local submission contains code, measurements, graphs and screenshots. It excludes access-key CSV files, '
         'private SSH keys and the encrypted temporary credential store. Billing should be rechecked after final usage posts.')

    output = DEST / 'Cloud_Resource_Management_Implementation.docx'
    # Remove inherited template rules, including Word's default title border.
    for node in document.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    for node in document.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    document.save(output)
    print(f'Created {output}; {figure_counter} figures. Render and inspect before delivery.')


if __name__ == '__main__':
    main()
